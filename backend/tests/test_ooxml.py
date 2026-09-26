"""O .xlsx gerado tem que abrir no Excel sem "reparar" -- nao so no openpyxl.

Bug real: o arquivo da etapa 8 passava em todos os testes (que releem com
openpyxl) e o Excel de verdade recusava e oferecia reparar, perdendo a
formatacao. Causa: Tabelas nomeadas com ref so no cabecalho (aba sem dados --
Avisos de quase todo documento, Itens de boleto) e celulas da Legenda tipadas
como texto sem texto. Ver tests/ooxml.py para as regras e por que o schema
sozinho nao basta.

Duas camadas:
- `ooxml.problemas` (sempre roda): invariantes do Excel lidos do XML bruto.
- validador oficial da Microsoft (opt-in, `--ooxml-sdk`, so Windows): schema
  OOXML via Open XML SDK. Complementar -- ele NAO pegou o bug acima.
"""
import os
import subprocess
import sys
import urllib.request
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ooxml  # noqa: E402

from app.excel_exporter import gerar_excel  # noqa: E402
from app.schemas import DocumentoExtraido, DocumentoParaExportar, ExtractionResult  # noqa: E402

CHAVE_44 = "35260472381189001001550010000123451123456786"


def _doc(arquivo="a.pdf", confiancas=None, avisos=None, corrigidos=None, modo="basico", **documento):
    documento.setdefault("tipo_documento", "nota_fiscal")
    return DocumentoParaExportar(
        arquivo=arquivo,
        resultado=ExtractionResult(
            modo_extracao=modo, confiancas=confiancas or {}, avisos=avisos or [], documento=DocumentoExtraido(**documento)
        ),
        corrigidos=corrigidos or {},
    )


# Um cenario por formato de exportacao que existe de verdade (e alguns de
# borda). O que importa e cobrir abas VAZIAS e CHEIAS, porque foi aba vazia que
# quebrou o arquivo.
CENARIOS = {
    "boleto (sem itens nem avisos)": [
        _doc(tipo_documento="boleto", numero_documento="123", data_vencimento="10/05/2026", valor_total=900.0,
             emissor="ESCOLA (CNPJ 11.111.111/0001-11)", destinatario="FULANO (CPF 000.000.000-00)",
             campos_adicionais=[{"campo": "Nosso Número", "valor": "10200000001-9"},
                                {"campo": "Valor do Documento", "valor": "900,00"}],
             confiancas={"numero_documento": "alta", "valor_total": "alta"})
    ],
    "danfe (itens, sem avisos)": [
        _doc(numero_documento="000012345", valor_total=229.0, data_emissao="15/04/2026",
             itens=[{"descricao": "Mochila", "quantidade": 1.0, "valor_unitario": 215.03, "valor_total": 215.03}],
             campos_adicionais=[{"campo": "Chave de Acesso", "valor": CHAVE_44}, {"campo": "CFOP", "valor": "5102"}],
             confiancas={"itens": "alta", "CFOP": "media", "valor_total": "alta"})
    ],
    "tudo preenchido, com media/baixa/corrigido": [
        _doc(numero_documento="1", emissor="E (CNPJ 11.222.333/0001-81)", valor_total=10.0, data_emissao="15/04/2026",
             itens=[{"descricao": "a", "valor_total": 1.0}, {"descricao": "b", "quantidade": "1 un", "valor_total": "x"}],
             campos_adicionais=[{"campo": "CFOP", "valor": "5102"}], avisos=["a soma nao bate"],
             confiancas={"emissor": "baixa", "data_emissao": "media", "itens": "media"},
             corrigidos={"valor_total": 9.0})
    ],
    "vazio (tipo desconhecido, nada extraido)": [_doc(tipo_documento="desconhecido")],
    "modo IA (sem confiancas)": [_doc(modo="ia", numero_documento="1", valor_total=1.0, itens=[{"descricao": "x"}])],
    "lote de 3": [
        _doc(arquivo="1.pdf", numero_documento="1", itens=[{"descricao": "a", "valor_total": 1.0}]),
        _doc(arquivo="2.pdf", tipo_documento="boleto", numero_documento="2", avisos=["x"]),
        _doc(arquivo="3.pdf", tipo_documento="desconhecido"),
    ],
    "lote de 5 tipos variados (com grafico, total, listas e regras)": [
        _doc(arquivo="nf1.pdf", numero_documento="000012345", valor_total=2290.0, data_emissao="15/04/2026",
             emissor="COMERCIAL ALFA LTDA (CNPJ 11.222.333/0001-81)", destinatario="FULANO DE TAL (CPF 000.000.000-00)",
             itens=[{"descricao": "Mochila", "quantidade": 2.0, "valor_unitario": 1145.0, "valor_total": 2290.0}],
             campos_adicionais=[{"campo": "CFOP", "valor": "5102"}, {"campo": "Chave de Acesso", "valor": CHAVE_44}],
             confiancas={"itens": "media", "CFOP": "media", "emissor": "alta"}),
        _doc(arquivo="bol1.pdf", tipo_documento="boleto", numero_documento="778", valor_total=900.0,
             data_vencimento="10/05/2026", emissor="ESCOLA BETA (CNPJ 22.333.444/0001-55)",
             campos_adicionais=[{"campo": "Nosso Número", "valor": "10200000001-9"}],
             confiancas={"valor_total": "alta", "Nosso Número": "baixa"}),
        _doc(arquivo="nf2.pdf", numero_documento="000054321", valor_total=15480.75, avisos=["a soma dos itens nao bate"],
             confiancas={"valor_total": "media"}, corrigidos={"numero_documento": "000054320"}),
        _doc(arquivo="ped.pdf", tipo_documento="pedido_compra", numero_documento="PC-77", valor_total=4320.0),
        _doc(arquivo="bol2.pdf", tipo_documento="boleto", numero_documento="779", valor_total=1150.0,
             data_vencimento="20/05/2026"),
    ],
    "lote com valor em texto (grafico sem esse documento, com uniao de intervalos)": [
        _doc(arquivo="1.pdf", numero_documento="1", valor_total=100.0),
        _doc(arquivo="2.pdf", numero_documento="2", valor_total="texto que nao virou numero"),
        _doc(arquivo="3.pdf", numero_documento="3", valor_total=50.0),
        _doc(arquivo="4.pdf", numero_documento="4", valor_total=70.0),
    ],
    "texto hostil (formula, controle, gigante)": [
        _doc(emissor='=HYPERLINK("http://x","y")', destinatario="A\x00B\x0bC", numero_documento="X" * 40000,
             itens=[{"descricao": "=SUM(A1:A2)"}], avisos=["=cmd|' /C calc'!A0"])
    ],
}


@pytest.mark.parametrize("nome", list(CENARIOS))
def test_arquivo_gerado_respeita_as_regras_do_excel(nome):
    achados = ooxml.problemas(gerar_excel(CENARIOS[nome], data_geracao=date(2026, 9, 23)))
    assert achados == [], f"{nome}: o Excel provavelmente pediria pra reparar:\n" + "\n".join(achados)


def test_aba_sem_dados_fica_sem_tabela_e_aba_com_dados_tem():
    wb = openpyxl.load_workbook(gerar_excel(CENARIOS["boleto (sem itens nem avisos)"]))
    assert wb["Itens"].tables == {} and wb["Avisos"].tables == {}
    assert "CamposAdicionais" in wb["Campos adicionais"].tables and "Documentos" in wb["Documentos"].tables


# ---------- o verificador pega o que tem que pegar (senao ele e decorativo) ----------


def _salvar(wb) -> bytes:
    from io import BytesIO

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_verificador_acusa_tabela_so_com_cabecalho():
    """Reproduz o bug que chegou ao Excel: tem que ser acusado."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID", "Aviso"])
    ws.add_table(Table(displayName="Avisos", ref="A1:B1"))
    assert any("sem nenhuma linha de dados" in a for a in ooxml.problemas(_salvar(wb)))


def test_verificador_acusa_celula_texto_sem_texto():
    wb = openpyxl.Workbook()
    wb.active["A1"] = ""
    wb.active["A1"].data_type = "s"
    # o openpyxl so escreve inlineStr vazio se a celula tiver estilo
    wb.active["A1"].font = openpyxl.styles.Font(bold=True)
    assert any("inlineStr sem <is>" in a for a in ooxml.problemas(_salvar(wb)))


def test_verificador_acusa_cabecalho_divergente_e_tabela_sobre_celula_mesclada():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID", "Nome"])
    ws.append([1, "a"])
    ws.append([2, "b"])
    tabela = Table(displayName="T", ref="A1:B3")
    ws.add_table(tabela)
    ws.merge_cells("A3:B3")
    conteudo = _salvar(wb)
    # adultera o nome da coluna no XML da tabela, sem mexer na celula
    from io import BytesIO

    entrada, saida = zipfile.ZipFile(BytesIO(conteudo)), BytesIO()
    with zipfile.ZipFile(saida, "w") as z:
        for item in entrada.infolist():
            dados = entrada.read(item.filename)
            if item.filename.startswith("xl/tables/"):
                dados = dados.replace(b'name="Nome"', b'name="Outro"')
            z.writestr(item, dados)
    achados = ooxml.problemas(saida.getvalue())
    assert any("difere de tableColumn" in a for a in achados)
    assert any("cruza celulas mescladas" in a for a in achados)


def _com_validacao(**kwargs):
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["A"])
    ws.append(["x"])
    dv = DataValidation(type="list", formula1=kwargs.pop("formula1", '"a,b"'), **kwargs)
    dv.add(kwargs.pop("faixa", "A2"))
    ws.add_data_validation(dv)
    return wb, ws


def test_verificador_aceita_validacao_de_dados_correta():
    wb, _ = _com_validacao(allow_blank=True, showErrorMessage=True, errorTitle="Valor inválido", error="Escolha da lista")
    assert ooxml.problemas(_salvar(wb)) == []


@pytest.mark.parametrize(
    "kwargs, trecho",
    [
        ({"formula1": '"' + ",".join(["item"] * 60) + '"'}, "lista literal"),  # 299 caracteres
        ({"showDropDown": True}, "ESCONDE a seta"),
        ({"errorTitle": "T" * 33}, "errorTitle"),
        ({"error": "E" * 256}, "error com"),
        ({"formula1": None}, "sem formula1"),
    ],
)
def test_verificador_acusa_validacao_de_dados_fora_dos_limites_do_excel(kwargs, trecho):
    wb, _ = _com_validacao(**kwargs)
    assert any(trecho in a for a in ooxml.problemas(_salvar(wb)))


def test_verificador_acusa_validacoes_sobrepostas():
    from openpyxl.worksheet.datavalidation import DataValidation

    wb, ws = _com_validacao()
    outra = DataValidation(type="list", formula1='"c,d"')
    outra.add("A2:A5")  # cobre A2, que ja tem a primeira validacao
    ws.add_data_validation(outra)
    assert any("sobrepoe outra validacao" in a for a in ooxml.problemas(_salvar(wb)))


def _adulterar(conteudo: bytes, parte_prefixo: str, de: bytes, para: bytes) -> bytes:
    from io import BytesIO

    entrada, saida = zipfile.ZipFile(BytesIO(conteudo)), BytesIO()
    with zipfile.ZipFile(saida, "w") as z:
        for item in entrada.infolist():
            dados = entrada.read(item.filename)
            if item.filename.startswith(parte_prefixo):
                dados = dados.replace(de, para)
            z.writestr(item, dados)
    return saida.getvalue()


def _com_regra_condicional() -> bytes:
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    ws.conditional_formatting.add(
        "A1:A3", CellIsRule(operator="equal", formula=['"x"'], fill=PatternFill("solid", start_color="FF0000", end_color="FF0000"))
    )
    return _salvar(wb)


def test_verificador_aceita_formatacao_condicional_correta():
    assert ooxml.problemas(_com_regra_condicional()) == []


@pytest.mark.parametrize(
    "parte, de, para, trecho",
    [
        ("xl/worksheets/sheet1.xml", b'dxfId="0"', b'dxfId="7"', "dxfId=7"),  # aponta pra dxf que nao existe
        ("xl/worksheets/sheet1.xml", b' operator="equal"', b"", "sem operator"),
        ("xl/worksheets/sheet1.xml", b' dxfId="0"', b"", "sem dxfId"),
        ("xl/worksheets/sheet1.xml", b'<formula>"x"</formula>', b"", "formula(s)"),
        ("xl/styles.xml", b'<dxfs count="1">', b'<dxfs count="4">', "dxfs count=4"),
    ],
)
def test_verificador_acusa_formatacao_condicional_invalida(parte, de, para, trecho):
    conteudo = _com_regra_condicional()
    adulterado = _adulterar(conteudo, parte, de, para)
    assert adulterado != conteudo, "a adulteracao nao encontrou o trecho (teste desatualizado)"
    assert any(trecho in a for a in ooxml.problemas(adulterado))


def _protegida(formula_desbloqueada=False, **protecao) -> bytes:
    from openpyxl.styles import Protection

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID", "Valor"])
    ws.append([1, 10.0])
    ws.append([2, 20.0])
    ws.add_table(Table(displayName="T", ref="A1:B3"))
    ws["B4"] = "=SUM(B2:B3)"
    if formula_desbloqueada:
        ws["B4"].protection = Protection(locked=False)
    ws.protection.sheet = True
    ws.protection.autoFilter = protecao.get("autoFilter", False)
    ws.protection.sort = protecao.get("sort", False)
    return _salvar(wb)


def test_verificador_aceita_aba_protegida_com_filtro_liberado_e_formula_bloqueada():
    assert ooxml.problemas(_protegida()) == []


def test_verificador_acusa_formula_desbloqueada_em_aba_protegida():
    assert any("desbloqueada numa aba protegida" in a for a in ooxml.problemas(_protegida(formula_desbloqueada=True)))


@pytest.mark.parametrize("atributo, trecho", [("autoFilter", "autoFilter bloqueado"), ("sort", "sort bloqueado")])
def test_verificador_acusa_aba_protegida_com_tabela_e_filtro_ou_ordenacao_bloqueados(atributo, trecho):
    assert any(trecho in a for a in ooxml.problemas(_protegida(**{atributo: True})))


def test_verificador_nao_cobra_nada_de_aba_sem_protecao():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID"])
    ws.append([1])
    ws.add_table(Table(displayName="T", ref="A1:A2"))
    ws["A3"] = "=SUM(A2:A2)"
    assert ooxml.problemas(_salvar(wb)) == []


def _com_grafico() -> bytes:
    return gerar_excel(CENARIOS["lote de 5 tipos variados (com grafico, total, listas e regras)"]).getvalue()


def test_verificador_aceita_o_grafico_gerado():
    assert ooxml.problemas(_com_grafico()) == []


@pytest.mark.parametrize(
    "parte, de, para, trecho",
    [
        # serie com menos categorias do que valores
        ("xl/charts/", b"$C$2:$C$6", b"$C$2:$C$5", "categoria(s)"),
        # referencia a uma aba que nao existe
        ("xl/charts/", b"<f>'Documentos'!$L$2:$L$6</f>", b"<f>'Inexistente'!$L$2:$L$6</f>", "aba inexistente"),
        # intervalo que o verificador nao entende
        ("xl/charts/", b"<f>'Documentos'!$L$2:$L$6</f>", b"<f>Documentos!L2..L6</f>", "nao reconhecida"),
        # eixo sem <delete val=0>: o bug do openpyxl 3.1 (eixo some no Excel 365)
        ("xl/charts/", b'<delete val="0" />', b"", "eixo"),
        # rotulo de valor com uma flag a menos
        ("xl/charts/", b'<showBubbleSize val="0" />', b"", "dLbls sem as flags"),
        # titulo sem overlay explicito
        ("xl/charts/", b'<overlay val="0" />', b"", "overlay"),
        # eixo que cruza um eixo que nao existe
        ("xl/charts/", b'<crossAx val="10" />', b'<crossAx val="77" />', "cruza um eixo inexistente"),
        # grafico ancorado dentro da tabela (linha 2), escondendo os dados
        # (com 5 documentos a ancora original e a linha 10 = <row>9</row>, 0-based)
        ("xl/drawings/drawing", b"<row>9</row>", b"<row>1</row>", "em cima de uma tabela"),
    ],
)
def test_verificador_acusa_grafico_quebrado(parte, de, para, trecho):
    conteudo = _com_grafico()
    adulterado = _adulterar(conteudo, parte, de, para)
    assert adulterado != conteudo, "a adulteracao nao encontrou o trecho (teste desatualizado)"
    assert any(trecho in a for a in ooxml.problemas(adulterado)), ooxml.problemas(adulterado)


# ---------- validador oficial da Microsoft (opt-in) ----------

_SDK_URL = "https://www.nuget.org/api/v2/package/DocumentFormat.OpenXml/2.20.0"
_SDK_DLL = "lib/net46/DocumentFormat.OpenXml.dll"

# Falso positivo documentado: o SDK exige ordem fixa nos filhos de <font>, mas o
# schema publicado define CT_Font como `choice maxOccurs=unbounded` (qualquer
# ordem) -- conferido em datypic.com/sc/ooxml/t-ssml_CT_Font.html. Aparece ate
# na fonte padrao que o openpyxl escreve em TODO arquivo, que o Excel abre sem
# reclamar. Qualquer OUTRO erro do SDK falha o teste.
_FALSO_POSITIVO_FONTE = ("/x:styleSheet[1]/x:fonts[1]/x:font[", "unexpected child element")


def _dll_do_sdk() -> Path:
    cache = Path(os.environ.get("LOCALAPPDATA", Path.home())) / "extrator-docs" / "openxml-sdk-2.20.0"
    dll = cache / _SDK_DLL
    if not dll.exists():
        cache.mkdir(parents=True, exist_ok=True)
        pacote = cache / "pacote.nupkg"
        urllib.request.urlretrieve(_SDK_URL, pacote)
        with zipfile.ZipFile(pacote) as z:
            z.extract(_SDK_DLL, cache)
    return dll


@pytest.mark.ooxml_sdk
@pytest.mark.parametrize("nome", list(CENARIOS))
def test_validador_oficial_da_microsoft_so_acusa_o_falso_positivo_conhecido(nome, tmp_path, request):
    if not request.config.getoption("--ooxml-sdk"):
        pytest.skip("validador da Microsoft e opt-in: pytest tests/test_ooxml.py --ooxml-sdk")
    if sys.platform != "win32":
        pytest.skip("usa o Windows PowerShell 5.1 (.NET Framework) pra carregar o SDK")

    arquivo = tmp_path / "saida.xlsx"
    arquivo.write_bytes(gerar_excel(CENARIOS[nome]).getvalue())
    script = Path(__file__).resolve().parent / "ooxml_sdk.ps1"
    r = subprocess.run(
        ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script),
         "-Dll", str(_dll_do_sdk()), "-Arquivo", str(arquivo)],
        capture_output=True, text=True, timeout=120,
    )
    assert r.returncode == 0, r.stderr
    erros = [linha.split("\t") for linha in r.stdout.splitlines() if linha.strip()]
    reais = [
        e for e in erros
        if not (e[1].startswith(_FALSO_POSITIVO_FONTE[0]) and _FALSO_POSITIVO_FONTE[1] in e[2])
    ]
    assert reais == [], "\n".join("\t".join(e) for e in reais)
