"""Excel multi-aba (etapa 6): estrutura, formatos, destaques e seguranca.

Le o .xlsx de volta com openpyxl -- o que o usuario abriria no Excel.
"""
import asyncio
import zipfile
from datetime import date, datetime
from io import BytesIO

import openpyxl
import pytest
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from app import main
from app.excel_exporter import (
    CABECALHOS_AVISOS,
    CABECALHOS_CAMPOS,
    CABECALHOS_ITENS,
    CABECALHOS_RESUMO,
    COR_AUSENTE,
    COR_BORDA,
    COR_CABECALHO_FUNDO,
    COR_CABECALHO_TEXTO,
    COR_DESTAQUE_VALOR_TOTAL,
    COR_ZEBRA,
    FUNDOS,
    FORMATO_DATA,
    FORMATO_MOEDA,
    FORMATO_MOEDA_UNITARIO,
    LARGURA_MAXIMA,
    LEGENDA_LINHAS,
    NOME_APLICATIVO,
    NOMES_TABELA,
    SEM_CONFIANCA,
    TAMANHO_FONTE_APLICATIVO,
    TAMANHO_FONTE_VALOR_PRINCIPAL,
    TAMANHO_FONTE_VALOR_TOTAL,
    VALORES_CONFIANCA,
    VALORES_TIPO,
    _lista_suspensa,
    gerar_excel,
    rotulo_documento,
    separar_documento_fiscal,
)
from app.schemas import (
    DocumentoExtraido,
    DocumentoParaExportar,
    ExportarExcelRequest,
    ExtractionResult,
)

CHAVE_44 = "35260472381189001001550010000123451123456786"
CHAVE_44_BLOCOS = " ".join(CHAVE_44[i:i + 4] for i in range(0, 44, 4))  # exibicao (Campos adicionais e Relatorio)

# cabecalhos "reais" de cada aba de DADOS -- usado pra distinguir a tabela de
# dados de celulas fora dela (nota de geracao em Documentos, linha de total
# nos Itens). A aba Relatorio NAO entra aqui: nao e uma tabela (linha 1 e o
# nome do aplicativo, nao um cabecalho de coluna) -- ver
# test_nenhuma_aba_de_dados_tem_linha_de_titulo_mesclada_acima_do_cabecalho.
_CABECALHOS_POR_ABA = {
    "Documentos": CABECALHOS_RESUMO,
    "Itens": CABECALHOS_ITENS,
    "Campos adicionais": CABECALHOS_CAMPOS,
    "Avisos": CABECALHOS_AVISOS,
}

# as 4 abas de DADOS (com Tabela nomeada, filtro, guia colorida, paisagem).
# Relatorio e a 1a aba (etapa 9), de LEITURA -- nao entra nesses contratos de
# proposito (retrato, sem Tabela, ver os testes especificos dela).
_ABAS_DE_DADOS = ["Documentos", "Itens", "Campos adicionais", "Avisos"]
_SHEETNAMES_ESPERADOS = ["Relatório", *_ABAS_DE_DADOS]


def _doc(arquivo="a.pdf", confiancas=None, avisos=None, aviso=None, corrigidos=None, **documento):
    documento.setdefault("tipo_documento", "nota_fiscal")
    return DocumentoParaExportar(
        arquivo=arquivo,
        resultado=ExtractionResult(
            modo_extracao="basico",
            confiancas=confiancas or {},
            avisos=avisos or [],
            aviso=aviso,
            documento=DocumentoExtraido(**documento),
        ),
        corrigidos=corrigidos or {},
    )


def _abrir(*docs):
    buffer = gerar_excel(list(docs))
    return openpyxl.load_workbook(buffer), buffer


def _linhas(ws):
    cabecalho = [c.value for c in ws[1]]
    return [dict(zip(cabecalho, linha)) for linha in ws.iter_rows(min_row=2, values_only=True)]


def _celula(ws, coluna, linha=2):
    cabecalho = [c.value for c in ws[1]]
    return ws.cell(row=linha, column=cabecalho.index(coluna) + 1)


def _fundo(celula):
    return celula.fill.fgColor.rgb[-6:] if celula.fill.fill_type == "solid" else None


def _linhas_sem_total(ws):
    """_linhas(), mas descartando a linha de total (Itens, ou Documentos num
    lote de 2+) -- e a unica linha da aba cujo ID nao e um inteiro."""
    return [linha for linha in _linhas(ws) if isinstance(linha["ID"], int)]


def _linha_com_texto(ws, texto, coluna=2, a_partir_de=1):
    """1a linha (a partir de `a_partir_de`) da aba Relatorio (rotulo/bloco
    livre, nao tabular) cuja celula da `coluna` tem exatamente esse texto.
    None se nao achar."""
    for linha in range(a_partir_de, ws.max_row + 1):
        if ws.cell(linha, coluna).value == texto:
            return linha
    return None


def _valor_por_rotulo(ws, rotulo, a_partir_de=1):
    """Celula de VALOR (coluna C, mesclada C:F) da 1a linha do Relatorio (a
    partir de `a_partir_de`) cujo rotulo (coluna B) e esse -- usado pra
    Emissor/Destinatario, onde os dois tem uma linha "Nome": passe a linha
    do titulo da secao certa pra nao pegar a do bloco anterior."""
    linha = _linha_com_texto(ws, rotulo, coluna=2, a_partir_de=a_partir_de)
    assert linha is not None, f"rotulo {rotulo!r} nao encontrado no Relatorio (a partir da linha {a_partir_de})"
    return ws.cell(linha, 3)


# ---------- estrutura ----------


def test_cinco_abas_sempre_presentes_na_ordem_relatorio_primeiro():
    wb, _ = _abrir(_doc(numero_documento="1"))
    assert wb.sheetnames == _SHEETNAMES_ESPERADOS  # Relatorio sempre primeiro (e a aba ativa)
    assert wb.active.title == "Relatório"
    for nome in ("Itens", "Campos adicionais", "Avisos"):
        assert wb[nome].max_row == 1, f"{nome} sem dados deve ter so o cabecalho"


def test_aba_sem_nenhum_dado_fica_oculta_mas_continua_existindo():
    """Estrutura estavel pra Power Query/formula (a aba existe, com o nome
    de sempre) sem mostrar aba vazia pra quem abre no Excel."""
    wb, _ = _abrir(_doc(numero_documento="1"))  # sem itens, campos ou avisos
    for nome in ("Itens", "Campos adicionais", "Avisos"):
        assert wb[nome].sheet_state == "hidden", f"{nome}: deveria ficar oculta quando vazia"
    assert wb["Relatório"].sheet_state == "visible" and wb["Documentos"].sheet_state == "visible"


def test_aba_com_dado_fica_visivel():
    wb, _ = _abrir(_doc(numero_documento="1", itens=[{"descricao": "x"}],
                        campos_adicionais=[{"campo": "CFOP", "valor": "5102"}], avisos=["a"]))
    for nome in ("Itens", "Campos adicionais", "Avisos"):
        assert wb[nome].sheet_state == "visible"


def test_cabecalho_negrito_congelado_e_com_tabela_nomeada_em_todas_as_abas_de_dados():
    wb, _ = _abrir(_doc(numero_documento="1", itens=[{"descricao": "x", "valor_total": 1.0}],
                        campos_adicionais=[{"campo": "CFOP", "valor": "5102"}], avisos=["a"]))
    for nome_aba in _ABAS_DE_DADOS:
        ws = wb[nome_aba]
        n = len(_CABECALHOS_POR_ABA[nome_aba])
        cabecalho = [ws.cell(1, c) for c in range(1, n + 1)]
        assert all(c.font.bold for c in cabecalho), f"{nome_aba}: cabecalho em negrito"
        # cabecalho E coluna ID congelados (a "primeira coluna" pedida)
        assert ws.freeze_panes == "B2", f"{nome_aba}: cabecalho e coluna ID congelados"
        # Itens com item(ns) ganha uma linha de total a mais, FORA da tabela
        ultima_linha = ws.max_row - 1 if nome_aba == "Itens" and ws.max_row > 1 else ws.max_row
        ultima_coluna = get_column_letter(n)
        nome_tabela = NOMES_TABELA[nome_aba]
        assert nome_tabela in ws.tables, f"{nome_aba}: sem Tabela nomeada"
        assert ws.tables[nome_tabela].ref == f"A1:{ultima_coluna}{ultima_linha}", f"{nome_aba}: tabela cobre so os dados"
        # nenhum auto_filter solto duplicado -- so a Tabela (ver docstring do modulo)
        assert ws.auto_filter.ref is None, f"{nome_aba}: auto_filter solto nao deveria existir (so a Tabela)"


def test_estilo_visual_proprio_da_tabela_desligado_pra_nao_brigar_com_a_zebra():
    wb, _ = _abrir(_doc(numero_documento="1"))
    tabela = wb["Documentos"].tables[NOMES_TABELA["Documentos"]]
    info = tabela.tableStyleInfo
    assert not any([info.showRowStripes, info.showColumnStripes, info.showFirstColumn, info.showLastColumn])


def test_toda_aba_de_dados_tem_id_sequencial_e_documento_legivel():
    wb, _ = _abrir(_doc(numero_documento="000012345", itens=[{"descricao": "x"}],
                        campos_adicionais=[{"campo": "CFOP", "valor": "5102"}], avisos=["a"]))
    for nome_aba in _ABAS_DE_DADOS:
        ws = wb[nome_aba]
        cabecalho = [c.value for c in ws[1]]
        assert cabecalho[:2] == ["ID", "Documento"] or (ws.title == "Documentos" and cabecalho[:3] == ["ID", "Arquivo", "Documento"])
        linha = _linhas(ws)[0]
        assert linha["ID"] == 1 and linha["Documento"] == "Nota fiscal 000012345"


def test_lote_de_dois_documentos_ids_documentos_e_linhas_ligados_pelo_id():
    d1 = _doc(arquivo="um.pdf", numero_documento="1", itens=[{"descricao": "a"}, {"descricao": "b"}],
              campos_adicionais=[{"campo": "CFOP", "valor": "5102"}], avisos=["aviso do 1"])
    d2 = _doc(arquivo="dois.pdf", tipo_documento="boleto", numero_documento="1",  # MESMO numero de proposito
              itens=[{"descricao": "c"}], campos_adicionais=[{"campo": "Parcela", "valor": "1/2"}])
    wb, _ = _abrir(d1, d2)

    resumo = _linhas_sem_total(wb["Documentos"])
    assert [(r["ID"], r["Arquivo"], r["Documento"]) for r in resumo] == [
        (1, "um.pdf", "Nota fiscal 1"), (2, "dois.pdf", "Boleto 1"),
    ]
    assert [(r["ID"], r["Descrição"]) for r in _linhas_sem_total(wb["Itens"])] == [(1, "a"), (1, "b"), (2, "c")]
    assert [(r["ID"], r["Campo"]) for r in _linhas(wb["Campos adicionais"])] == [(1, "CFOP"), (2, "Parcela")]
    assert [(r["ID"], r["Aviso"]) for r in _linhas(wb["Avisos"])] == [(1, "aviso do 1")]


def test_rotulo_documento_com_e_sem_numero():
    assert rotulo_documento("boleto", "123", "a.pdf") == "Boleto 123"
    assert rotulo_documento("nota_fiscal", None, "a.pdf") == "Nota fiscal — a.pdf"
    assert rotulo_documento("nota_fiscal", "  ", None) == "Nota fiscal (sem número)"
    assert rotulo_documento("tipo_inventado", "9", None) == "tipo_inventado 9"


# ---------- Documentos ----------


def test_emissor_e_destinatario_sem_o_documento_entre_parenteses_e_em_colunas_proprias():
    wb, _ = _abrir(_doc(emissor="DELL LTDA (CNPJ 72.381.189/0010-01)", destinatario="FULANO DE TAL (CPF 000.000.000-00)"))
    r = _linhas(wb["Documentos"])[0]
    assert r["Emissor"] == "DELL LTDA" and r["Emissor CNPJ/CPF"] == "72.381.189/0010-01"
    assert r["Destinatário"] == "FULANO DE TAL" and r["Destinatário CNPJ/CPF"] == "000.000.000-00"


@pytest.mark.parametrize(
    "entrada, esperado",
    [
        ("EMPRESA SEM DOC LTDA", ("EMPRESA SEM DOC LTDA", None)),
        ("EMPRESA (FILIAL) LTDA (CNPJ 11.222.333/0001-81)", ("EMPRESA (FILIAL) LTDA", "11.222.333/0001-81")),
        ("(CNPJ 11.222.333/0001-81)", ("(CNPJ 11.222.333/0001-81)", None)),  # sem nome: nao inventa
        (None, (None, None)),
        ("  ", (None, None)),
    ],
)
def test_separar_documento_fiscal(entrada, esperado):
    assert separar_documento_fiscal(entrada) == esperado


def test_numero_do_documento_fica_texto_e_preserva_zeros_a_esquerda():
    wb, _ = _abrir(_doc(numero_documento="000012345"))
    celula = _celula(wb["Documentos"], "Número")
    assert celula.value == "000012345" and celula.data_type == "s"


def test_datas_viram_data_de_verdade_e_o_que_nao_converte_fica_texto():
    wb, _ = _abrir(
        _doc(data_emissao="15/04/2026", data_vencimento="2026-05-01"),
        _doc(data_emissao="31/02/2026", data_vencimento="quando der"),
    )
    ws = wb["Documentos"]
    emissao, vencimento = _celula(ws, "Data de emissão", 2), _celula(ws, "Data de vencimento", 2)
    assert isinstance(emissao.value, datetime) and emissao.value.date() == date(2026, 4, 15)
    assert emissao.number_format == FORMATO_DATA and emissao.is_date
    assert vencimento.value.date() == date(2026, 5, 1)  # ISO (o modo IA pode devolver)
    # dia inexistente e texto livre: permanecem texto, sem quebrar
    assert _celula(ws, "Data de emissão", 3).value == "31/02/2026"
    assert _celula(ws, "Data de vencimento", 3).value == "quando der"


def test_valor_total_float_vira_numero_em_moeda_e_string_fica_texto():
    wb, _ = _abrir(_doc(valor_total=1234.5), _doc(valor_total="a combinar"))
    ws = wb["Documentos"]
    numero = _celula(ws, "Valor total", 2)
    assert numero.value == 1234.5 and numero.number_format == FORMATO_MOEDA and numero.data_type == "n"
    assert _celula(ws, "Valor total", 3).value == "a combinar"


def test_formato_monetario_usa_a_notacao_interna_que_o_excel_localiza():
    """O codigo guardado no arquivo e en-US; em pt-BR o Excel mostra R$ 1.234,56.
    Escrever "#.##0,00" literalmente quebraria o formato."""
    assert FORMATO_MOEDA == '"R$" #,##0.00'
    assert "#.##0,00" not in FORMATO_MOEDA and "#.##0,00" not in FORMATO_MOEDA_UNITARIO


def test_tipo_no_resumo_e_o_valor_interno_e_a_confianca_geral_segue_a_regra_da_tela():
    d = _doc(numero_documento="1", emissor="E (CNPJ 1)", valor_total=10.0,
             campos_adicionais=[{"campo": "CFOP", "valor": "5102"}],
             confiancas={"numero_documento": "alta", "emissor": "alta", "valor_total": "media", "CFOP": "media"})
    r = _linhas(_abrir(d)[0]["Documentos"])[0]
    assert r["Tipo"] == "nota_fiscal"  # interno (bom pra filtrar), nao "Nota fiscal"
    assert r["Confiança geral"] == "2 de 4 alta"


def test_confianca_geral_no_modo_ia_e_traco():
    r = _linhas(_abrir(_doc(numero_documento="1"))[0]["Documentos"])[0]
    assert r["Confiança geral"] == "—"


def test_confianca_geral_deixa_os_corrigidos_fora_do_x_de_y():
    d = _doc(numero_documento="1", emissor="E", valor_total=10.0,
             confiancas={"numero_documento": "alta", "emissor": "alta", "valor_total": "alta"},
             corrigidos={"emissor": "outro"})
    assert _linhas(_abrir(d)[0]["Documentos"])[0]["Confiança geral"] == "2 de 2 alta"


# ---------- Campos adicionais ----------


def test_campos_monetarios_viram_numero_e_codigos_ficam_texto():
    d = _doc(campos_adicionais=[
        {"campo": "Valor do Documento", "valor": "1.000,00"},
        {"campo": "Desconto", "valor": "100,00"},
        {"campo": "Valor a Pagar", "valor": "sem valor"},  # nao parseavel: texto
        {"campo": "Chave de Acesso", "valor": CHAVE_44},
        {"campo": "Nosso Número", "valor": "10200000001-9"},
        {"campo": "CFOP", "valor": "5102"},
    ])
    ws = _abrir(d)[0]["Campos adicionais"]
    por_campo = {ws.cell(row=r, column=3).value: ws.cell(row=r, column=4) for r in range(2, ws.max_row + 1)}
    assert por_campo["Valor do Documento"].value == 1000.0 and por_campo["Valor do Documento"].number_format == FORMATO_MOEDA
    assert por_campo["Desconto"].value == 100.0
    assert por_campo["Valor a Pagar"].value == "sem valor"
    # 44 digitos: como numero o Excel guardaria so 15 de precisao -- tem que ser texto INTEIRO.
    # Em blocos de 4 (mesma exibicao da tela e do Relatorio, ver _formatar_chave_acesso).
    assert por_campo["Chave de Acesso"].value == CHAVE_44_BLOCOS and por_campo["Chave de Acesso"].data_type == "s"
    assert por_campo["Nosso Número"].value == "10200000001-9"
    assert por_campo["CFOP"].value == "5102" and por_campo["CFOP"].data_type == "s"


def test_coluna_confianca_dos_campos_adicionais_em_texto():
    d = _doc(campos_adicionais=[{"campo": "A", "valor": "1"}, {"campo": "B", "valor": "2"},
                                {"campo": "C", "valor": "3"}, {"campo": "D", "valor": "4"}],
             confiancas={"A": "alta", "B": "media", "C": "baixa"}, corrigidos={"D": "x"})
    linhas = _linhas(_abrir(d)[0]["Campos adicionais"])
    assert [(l["Campo"], l["Confiança"]) for l in linhas] == [
        ("A", "Alta"), ("B", "Média"), ("C", "Baixa"), ("D", "Corrigido"),
    ]


# ---------- protecao de planilha (sem senha) ----------


def _lote_para_protecao():
    return [
        _doc(arquivo="1.pdf", numero_documento="1", valor_total=100.0, campos_adicionais=[{"campo": "CFOP", "valor": "5102"}],
             itens=[{"descricao": "a", "quantidade": 1.0, "valor_total": 100.0}], avisos=["x"]),
        _doc(arquivo="2.pdf", tipo_documento="boleto", numero_documento="2", valor_total=50.0),
    ]


def test_abas_de_dados_protegidas_sem_senha_e_relatorio_livre():
    wb, _ = _abrir(*_lote_para_protecao())
    for nome in _ABAS_DE_DADOS:
        p = wb[nome].protection
        assert p.sheet is True, nome
        assert not p.password and not p.hashValue, f"{nome}: a trava e so contra engano, sem senha"
    assert wb["Relatório"].protection.sheet is False


def test_protecao_libera_filtro_ordenacao_e_largura_de_coluna():
    """Semantica invertida do formato: False = LIBERADO. Filtro funcionando vale
    mais que a trava (decisao do projeto)."""
    for nome in _ABAS_DE_DADOS:
        p = _abrir(*_lote_para_protecao())[0][nome].protection
        assert p.autoFilter is False and p.sort is False and p.formatColumns is False and p.formatRows is False, nome
        assert p.objects is False  # grafico continua selecionavel/copiavel


def test_cabecalho_e_linhas_de_total_bloqueados_e_celulas_de_dado_livres():
    wb, _ = _abrir(*_lote_para_protecao())
    for nome in _ABAS_DE_DADOS:
        ws = wb[nome]
        n_colunas = len(_CABECALHOS_POR_ABA[nome])
        assert all(ws.cell(1, c).protection.locked for c in range(1, n_colunas + 1)), f"{nome}: cabecalho tem que ser bloqueado"
    # dados: 2 documentos em Documentos; 1 item; 1 campo; 1 aviso
    for nome, n_linhas in (("Documentos", 2), ("Itens", 1), ("Campos adicionais", 1), ("Avisos", 1)):
        ws = wb[nome]
        for linha in range(2, n_linhas + 2):
            for c in range(1, len(_CABECALHOS_POR_ABA[nome]) + 1):
                assert not ws.cell(linha, c).protection.locked, f"{nome} {ws.cell(linha, c).coordinate} deveria ser livre"
        for c in range(1, len(_CABECALHOS_POR_ABA[nome]) + 1):  # a linha logo abaixo (total, quando ha) e bloqueada
            assert ws.cell(n_linhas + 2, c).protection.locked, f"{nome} {ws.cell(n_linhas + 2, c).coordinate} deveria ser bloqueada"


def test_nenhuma_formula_do_arquivo_fica_desbloqueada():
    """A razao de ser da trava: o =SUM dos totais (Itens e Documentos)."""
    wb, _ = _abrir(*_lote_para_protecao())
    formulas = [c for nome in wb.sheetnames for linha in wb[nome].iter_rows() for c in linha if c.data_type == "f"]
    assert len(formulas) >= 3  # Itens: 2 (quantidade e valor); Documentos: 1 (total do lote)
    assert all(c.protection.locked for c in formulas), [c.coordinate for c in formulas if not c.protection.locked]


# ---------- grafico do lote (aba Documentos) ----------

_NS_GRAFICO = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
               "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"}


def _lote_com_valores(*valores):
    return [_doc(arquivo=f"{i}.pdf", numero_documento=str(i), valor_total=v) for i, v in enumerate(valores, start=1)]


def _xml(buffer, parte):
    from xml.etree import ElementTree as ET

    with zipfile.ZipFile(BytesIO(buffer.getvalue())) as z:
        return ET.fromstring(z.read(parte)) if parte in z.namelist() else None


def _formulas_da_serie(buffer):
    grafico = _xml(buffer, "xl/charts/chart1.xml")
    serie = grafico.find(".//c:ser", _NS_GRAFICO)
    return {papel: serie.find(f"c:{papel}//c:f", _NS_GRAFICO).text for papel in ("tx", "cat", "val")}


def test_lote_com_valores_numericos_gera_um_grafico_de_barras_de_valor_por_documento():
    buffer = gerar_excel(_lote_com_valores(100.0, 50.5, 80.0))
    letra_nome = get_column_letter(CABECALHOS_RESUMO.index("Documento") + 1)
    letra_valor = get_column_letter(CABECALHOS_RESUMO.index("Valor total") + 1)
    assert _formulas_da_serie(buffer) == {
        "tx": f"'Documentos'!${letra_valor}$1",
        "cat": f"'Documentos'!${letra_nome}$2:${letra_nome}$4",
        "val": f"'Documentos'!${letra_valor}$2:${letra_valor}$4",
    }
    grafico = _xml(buffer, "xl/charts/chart1.xml")
    assert grafico.find(".//c:barDir", _NS_GRAFICO).get("val") == "col"
    # o que o grafico referencia e mesmo o que ele diz que e (nao so letras "certas")
    ws = openpyxl.load_workbook(buffer)["Documentos"]
    assert ws[f"{letra_valor}1"].value == "Valor total" and ws[f"{letra_nome}1"].value == "Documento"
    assert [ws[f"{letra_valor}{r}"].value for r in (2, 3, 4)] == [100.0, 50.5, 80.0]
    assert [ws[f"{letra_nome}{r}"].value for r in (2, 3, 4)] == ["Nota fiscal 1", "Nota fiscal 2", "Nota fiscal 3"]


def test_grafico_nao_aparece_com_um_unico_documento():
    buffer = gerar_excel(_lote_com_valores(100.0))
    with zipfile.ZipFile(BytesIO(buffer.getvalue())) as z:
        assert not [n for n in z.namelist() if n.startswith(("xl/charts/", "xl/drawings/drawing"))]


def test_grafico_nao_aparece_com_dois_documentos_se_so_um_tem_valor_numerico():
    """1 barra so seria enfeite -- vale o mesmo criterio do "2 ou mais"."""
    buffer = gerar_excel(_lote_com_valores(100.0, "texto", None))
    assert _xml(buffer, "xl/charts/chart1.xml") is None


def test_grafico_deixa_de_fora_o_documento_cujo_valor_total_e_texto():
    """O Excel plota texto como ZERO: incluir a linha mostraria uma barra
    "R$ 0,00" que nao existe. Fica de fora, via uniao de intervalos."""
    buffer = gerar_excel(_lote_com_valores(100.0, "texto que nao virou numero", 50.0, 70.0))
    f = _formulas_da_serie(buffer)
    assert f["val"] == "('Documentos'!$L$2,'Documentos'!$L$4:$L$5)"
    assert f["cat"] == "('Documentos'!$C$2,'Documentos'!$C$4:$C$5)"  # mesmas linhas: categoria e valor alinhados


def test_referencia_de_linhas_contigua_unitaria_com_lacuna_e_nome_de_aba_com_apostrofo():
    from app.excel_exporter import _referencia_de_linhas

    assert _referencia_de_linhas("Documentos", 12, [2, 3, 4]) == "'Documentos'!$L$2:$L$4"
    assert _referencia_de_linhas("Documentos", 12, [5]) == "'Documentos'!$L$5"
    assert _referencia_de_linhas("Documentos", 12, [2, 3, 6, 8, 9]) == "('Documentos'!$L$2:$L$3,'Documentos'!$L$6,'Documentos'!$L$8:$L$9)"
    assert _referencia_de_linhas("Dado's", 1, [2]) == "'Dado''s'!$A$2"


def test_grafico_fica_abaixo_da_linha_de_total_na_coluna_b_e_fora_da_tabela():
    n = 4
    buffer = gerar_excel(_lote_com_valores(*[10.0 * i for i in range(1, n + 1)]))
    ancora = _xml(buffer, "xl/drawings/drawing1.xml").find("xdr:oneCellAnchor/xdr:from", _NS_GRAFICO)
    col, linha = int(ancora.find("xdr:col", _NS_GRAFICO).text) + 1, int(ancora.find("xdr:row", _NS_GRAFICO).text) + 1
    linha_total = n + 2
    assert col == 2  # B: a coluna A (ID) e congelada, o grafico nao pode atravessar a divisa
    assert linha == linha_total + 3  # 2 linhas em branco depois da linha de total
    ws = openpyxl.load_workbook(buffer)["Documentos"]
    assert ws.tables[NOMES_TABELA["Documentos"]].ref.endswith(f"{n + 1}")  # a tabela acaba antes do grafico


def test_grafico_com_eixos_explicitos_sem_legenda_e_com_rotulo_de_valor():
    grafico = _xml(gerar_excel(_lote_com_valores(100.0, 50.0)), "xl/charts/chart1.xml")
    for eixo in ("catAx", "valAx"):
        # openpyxl 3.1 nao escreve <delete>; sem ele o Excel 365 esconde o eixo
        assert grafico.find(f".//c:{eixo}/c:delete", _NS_GRAFICO).get("val") == "0"
    assert grafico.find(".//c:legend", _NS_GRAFICO) is None  # 1 serie so
    assert grafico.find(".//c:ser/c:cat/c:strRef", _NS_GRAFICO) is not None  # nomes, nao numeros
    assert grafico.find(".//c:ser/c:cat/c:numRef", _NS_GRAFICO) is None
    rotulos = grafico.find(".//c:ser/c:dLbls", _NS_GRAFICO)
    assert rotulos.find("c:showVal", _NS_GRAFICO).get("val") == "1"
    for flag in ("showLegendKey", "showCatName", "showSerName", "showPercent", "showBubbleSize"):
        assert rotulos.find(f"c:{flag}", _NS_GRAFICO).get("val") == "0", flag  # as omitidas o Excel pode ligar
    assert grafico.find(".//c:title/c:overlay", _NS_GRAFICO).get("val") == "0"
    assert "Valor total por documento" in "".join(grafico.find(".//c:title", _NS_GRAFICO).itertext())


def test_grafico_fica_fora_da_area_de_impressao():
    ws = _abrir(*_lote_com_valores(100.0, 50.0, 80.0))[0]["Documentos"]
    assert ws.print_area.endswith("$5")  # cabecalho + 3 documentos + linha de total; o grafico (a partir da 8) nao entra


# ---------- validacao de dados (listas suspensas) e formatacao condicional ----------


def _lote_cobrindo_todos_os_estados():
    """5 documentos, um de cada tipo interno; o 1o tem 5 campos adicionais,
    um em cada estado de confianca (alta/media/baixa/corrigido/sem info)."""
    campos = [{"campo": c, "valor": str(i)} for i, c in enumerate("ABCDE", start=1)]
    return [
        _doc(arquivo="1.pdf", tipo_documento="nota_fiscal", numero_documento="1", campos_adicionais=campos,
             confiancas={"A": "alta", "B": "media", "C": "baixa"}, corrigidos={"D": "x"}),  # E: sem confianca
        _doc(arquivo="2.pdf", tipo_documento="boleto", numero_documento="2"),
        _doc(arquivo="3.pdf", tipo_documento="pedido_compra", numero_documento="3"),
        _doc(arquivo="4.pdf", tipo_documento="relatorio", numero_documento="4"),
        _doc(arquivo="5.pdf", tipo_documento="desconhecido"),
    ]


def _validacoes(ws):
    return ws.data_validations.dataValidation


def _regras(ws):
    """[(sqref, formula, regra)] da formatacao condicional da aba."""
    return [(str(cf.sqref), r.formula[0], r) for cf in ws.conditional_formatting for r in cf.rules]


def test_tipo_tem_lista_suspensa_com_os_valores_internos_so_nas_linhas_de_dados():
    ws = _abrir(*_lote_cobrindo_todos_os_estados())[0]["Documentos"]
    (dv,) = _validacoes(ws)
    assert dv.type == "list" and str(dv.sqref) == "D2:D6"  # 5 documentos; a linha 7 (total do lote) fica de fora
    assert dv.formula1 == '"boleto,nota_fiscal,pedido_compra,relatorio,desconhecido"'  # internos, nao os rotulos amigaveis
    assert VALORES_TIPO == ["boleto", "nota_fiscal", "pedido_compra", "relatorio", "desconhecido"]
    assert dv.showDropDown in (None, False)  # True ESCONDERIA a seta (semantica invertida do formato)
    assert dv.showErrorMessage and dv.errorStyle == "stop" and dv.allow_blank
    assert "nota_fiscal" in dv.error  # a mensagem diz quais valores valem


def test_confianca_dos_campos_adicionais_tem_lista_suspensa_com_o_dominio_real_da_coluna():
    ws = _abrir(*_lote_cobrindo_todos_os_estados())[0]["Campos adicionais"]
    (dv,) = _validacoes(ws)
    assert dv.type == "list" and str(dv.sqref) == "E2:E6"
    assert dv.formula1 == '"Alta,Média,Baixa,Corrigido,—"'
    assert VALORES_CONFIANCA == ["Alta", "Média", "Baixa", "Corrigido", SEM_CONFIANCA]
    assert dv.showDropDown in (None, False) and dv.showErrorMessage and dv.errorStyle == "stop"


def test_todo_valor_que_o_exportador_escreve_nas_colunas_validadas_pertence_a_propria_lista():
    """O arquivo nao pode violar a regra que ele mesmo carrega ("Circle invalid
    data" do Excel marcaria o proprio conteudo gerado)."""
    wb, _ = _abrir(*_lote_cobrindo_todos_os_estados())
    tipos = [r["Tipo"] for r in _linhas_sem_total(wb["Documentos"])]
    assert set(tipos) == set(VALORES_TIPO)  # o lote cobre os 5, e todos estao na lista
    confiancas = [r["Confiança"] for r in _linhas(wb["Campos adicionais"])]
    assert set(confiancas) == set(VALORES_CONFIANCA)  # idem: alta/media/baixa/corrigido/sem info


def test_so_tipo_e_confianca_dos_campos_ganham_lista_suspensa():
    """"Confiança geral" (Documentos) e uma proporcao em texto ("6 de 7 alta"),
    nao um conjunto fechado -- lista suspensa ali seria errada."""
    wb, _ = _abrir(*_lote_cobrindo_todos_os_estados())
    assert [str(dv.sqref) for dv in _validacoes(wb["Documentos"])] == ["D2:D6"]
    assert [str(dv.sqref) for dv in _validacoes(wb["Campos adicionais"])] == ["E2:E6"]
    for nome in ("Relatório", "Itens", "Avisos"):
        assert _validacoes(wb[nome]) == [], nome


def test_aba_sem_dados_nao_tem_validacao_nem_regra_condicional():
    wb, _ = _abrir(_doc(tipo_documento="boleto", numero_documento="1"))  # sem campos adicionais
    ws = wb["Campos adicionais"]  # oculta: nao ha linha nenhuma pra validar/colorir
    assert _validacoes(ws) == [] and _regras(ws) == []


def test_lista_suspensa_recusa_o_que_o_excel_recusaria():
    ws = openpyxl.Workbook().active
    with pytest.raises(ValueError, match="virgula"):
        _lista_suspensa(ws, "A2:A3", ["a,b", "c"])  # a virgula separaria em 3 itens
    with pytest.raises(ValueError, match="limite"):
        _lista_suspensa(ws, "A2:A3", ["x" * 100] * 3)  # 302 caracteres > 255


def test_coluna_confianca_tem_regras_condicionais_nativas_por_texto():
    ws = _abrir(*_lote_cobrindo_todos_os_estados())[0]["Campos adicionais"]
    regras = _regras(ws)
    assert {faixa for faixa, _, _ in regras} == {"E2:E6"}
    por_texto = {formula: regra for _, formula, regra in regras}
    assert set(por_texto) == {'"Média"', '"Baixa"', '"Corrigido"'}  # "Alta" e "—" nao pintam
    assert all(r.type == "cellIs" and r.operator == "equal" for r in por_texto.values())
    assert por_texto['"Média"'].dxf.fill.bgColor.rgb[-6:] == FUNDOS["media"]
    assert por_texto['"Baixa"'].dxf.fill.bgColor.rgb[-6:] == FUNDOS["baixa"]
    assert por_texto['"Corrigido"'].dxf.fill.bgColor.rgb[-6:] == FUNDOS["corrigido"]
    assert por_texto['"Corrigido"'].dxf.font.i  # "corrigido" tambem em italico, como no fundo fixo
    assert not por_texto['"Média"'].dxf.font and not por_texto['"Baixa"'].dxf.font


def test_as_regras_condicionais_casam_exatamente_com_o_texto_que_o_exportador_escreve():
    """Simula o Excel: pra cada celula da coluna, qual regra dispara. Pega o
    erro silencioso de literal diferente do texto ("Media" sem acento na regra
    contra "Média" na celula: a cor nunca apareceria e nada quebraria)."""
    ws = _abrir(*_lote_cobrindo_todos_os_estados())[0]["Campos adicionais"]
    regras = _regras(ws)

    def cor_que_o_excel_aplicaria(texto):
        for _, formula, regra in regras:
            if formula.strip('"') == texto:
                return regra.dxf.fill.bgColor.rgb[-6:]
        return None

    coluna = CABECALHOS_CAMPOS.index("Confiança") + 1
    textos = [ws.cell(linha, coluna).value for linha in range(2, 7)]
    assert textos == ["Alta", "Média", "Baixa", "Corrigido", SEM_CONFIANCA]
    assert [cor_que_o_excel_aplicaria(t) for t in textos] == [None, "FDF3E0", "FBECEB", "E7F0EF", None]


def test_celula_de_confianca_nao_tem_fundo_fixo_a_cor_vem_da_regra():
    """Se sobrasse um fundo fixo, editar "Baixa" -> "Alta" deixaria a celula
    vermelha (a regra deixa de disparar e o fundo antigo reaparece)."""
    ws = _abrir(*_lote_cobrindo_todos_os_estados())[0]["Campos adicionais"]
    coluna = CABECALHOS_CAMPOS.index("Confiança") + 1
    fundos = {_fundo(ws.cell(linha, coluna)) for linha in range(2, 7)}
    assert fundos <= {None, COR_ZEBRA}, fundos  # so a zebra (que a regra, quando dispara, sobrepoe)


def test_valor_ao_lado_da_confianca_mantem_fundo_fixo_e_comentario_de_proveniencia():
    """A cor do CAMPO documenta como ele foi extraido; reescrever o valor nao
    pode apagar esse rastro. So a coluna Confianca (cujo texto e o proprio
    valor) virou regra condicional."""
    ws = _abrir(*_lote_cobrindo_todos_os_estados())[0]["Campos adicionais"]
    coluna = CABECALHOS_CAMPOS.index("Valor") + 1
    assert _fundo(ws.cell(3, coluna)) == FUNDOS["media"] and "média" in ws.cell(3, coluna).comment.text
    assert _fundo(ws.cell(4, coluna)) == FUNDOS["baixa"]
    assert _fundo(ws.cell(5, coluna)) == FUNDOS["corrigido"] and ws.cell(5, coluna).font.italic


# ---------- Itens ----------


def test_itens_numericos_com_formato_e_texto_quando_nao_converteu():
    d = _doc(itens=[{"descricao": "Mochila", "quantidade": 2.5, "valor_unitario": 215.0312, "valor_total": 537.58},
                    {"descricao": "Servico", "quantidade": "1 un", "valor_unitario": None, "valor_total": "cortesia"}])
    ws = _abrir(d)[0]["Itens"]
    assert _celula(ws, "Quantidade", 2).value == 2.5
    unit = _celula(ws, "Valor unitário", 2)
    assert unit.value == 215.0312 and unit.number_format == FORMATO_MOEDA_UNITARIO  # ate 4 casas
    assert _celula(ws, "Valor total", 2).number_format == FORMATO_MOEDA
    assert _celula(ws, "Quantidade", 3).value == "1 un" and _celula(ws, "Valor total", 3).value == "cortesia"


# ---------- Avisos ----------


def test_avisos_uma_linha_por_aviso_e_cai_na_string_aviso_quando_nao_ha_lista():
    wb, _ = _abrir(_doc(numero_documento="1", avisos=["primeiro", "segundo"]),
                   _doc(numero_documento="2", aviso="so a string antiga"))
    assert [(l["ID"], l["Aviso"]) for l in _linhas(wb["Avisos"])] == [
        (1, "primeiro"), (1, "segundo"), (2, "so a string antiga"),
    ]


# ---------- destaques ----------


def test_fundo_e_comentario_para_media_baixa_e_corrigido():
    d = _doc(numero_documento="1", data_emissao="15/04/2026", emissor="E (CNPJ 11.222.333/0001-81)",
             valor_total=10.0,
             confiancas={"numero_documento": "alta", "data_emissao": "media", "emissor": "baixa", "valor_total": "alta"},
             corrigidos={"valor_total": 9.0})
    ws = _abrir(d)[0]["Documentos"]

    assert _fundo(_celula(ws, "Número")) is None and _celula(ws, "Número").comment is None  # alta: sem destaque
    data = _celula(ws, "Data de emissão")
    assert _fundo(data) == "FDF3E0" and "média" in data.comment.text
    emissor, emissor_doc = _celula(ws, "Emissor"), _celula(ws, "Emissor CNPJ/CPF")
    assert _fundo(emissor) == "FBECEB" and "baixa" in emissor.comment.text
    assert _fundo(emissor_doc) == "FBECEB" and emissor_doc.comment is None  # CNPJ leva a cor, o comentario fica no nome
    total = _celula(ws, "Valor total")
    assert _fundo(total) == "E7F0EF" and total.font.italic  # corrigido: outra cor + italico
    assert total.comment.text == "Corrigido pelo usuário. Valor extraído: 9,00"


def test_corrigido_para_vazio_ainda_e_marcado_com_o_valor_original():
    d = _doc(numero_documento=None, corrigidos={"numero_documento": "000012345"})
    celula = _celula(_abrir(d)[0]["Documentos"], "Número")
    assert celula.value is None and _fundo(celula) == "E7F0EF"
    assert "000012345" in celula.comment.text


def test_confianca_media_da_tabela_de_itens_pinta_as_linhas_com_um_comentario_so():
    d = _doc(itens=[{"descricao": "a", "valor_total": 1.0}, {"descricao": "b", "valor_total": 2.0}],
             confiancas={"itens": "media"})
    ws = _abrir(d)[0]["Itens"]
    assert _fundo(_celula(ws, "Descrição", 2)) == "FDF3E0" and _fundo(_celula(ws, "Valor total", 3)) == "FDF3E0"
    assert _celula(ws, "Descrição", 2).comment is not None
    assert _celula(ws, "Descrição", 3).comment is None


def test_modo_ia_nao_pinta_nada():
    ws = _abrir(_doc(numero_documento="1", valor_total=1.0))[0]["Documentos"]
    assert all(_fundo(c) is None for c in ws[2])


# ---------- identidade visual (etapa 7) ----------


def test_cabecalho_fundo_escuro_texto_branco_negrito_e_linha_mais_alta():
    wb, _ = _abrir(_doc(numero_documento="1"))
    ws = wb["Documentos"]
    cabecalho = ws.cell(1, 1)
    assert _fundo(cabecalho) == COR_CABECALHO_FUNDO
    assert cabecalho.font.color.rgb[-6:] == COR_CABECALHO_TEXTO
    assert cabecalho.font.bold
    assert ws.row_dimensions[1].height > 15  # padrao do openpyxl e ~15


def test_bordas_finas_em_toda_celula_cabecalho_e_dado():
    wb, _ = _abrir(_doc(numero_documento="1"))
    ws = wb["Documentos"]
    for celula in (ws.cell(1, 1), ws.cell(2, 1)):
        for lado in (celula.border.left, celula.border.right, celula.border.top, celula.border.bottom):
            assert lado.style == "thin" and lado.color.rgb[-6:] == COR_BORDA


def test_alinhamento_por_tipo_texto_esquerda_numero_e_data_a_direita_com_recuo():
    wb, _ = _abrir(_doc(numero_documento="000123", data_emissao="15/04/2026", valor_total=10.0))
    ws = wb["Documentos"]
    texto, numero, data = _celula(ws, "Número"), _celula(ws, "Valor total"), _celula(ws, "Data de emissão")
    assert texto.alignment.horizontal == "left" and texto.alignment.indent == 1
    assert numero.alignment.horizontal == "right" and numero.alignment.indent == 1
    assert data.alignment.horizontal == "right" and data.alignment.indent == 1


def test_zebra_alterna_por_linha_e_destaque_de_confianca_sempre_vence():
    d1 = _doc(numero_documento="1", emissor="A")
    d2 = _doc(numero_documento="2", emissor="B", confiancas={"emissor": "baixa"})
    d3 = _doc(numero_documento="3", emissor="C")
    d4 = _doc(numero_documento="4", emissor="D")
    wb, _ = _abrir(d1, d2, d3, d4)
    ws = wb["Documentos"]
    assert _fundo(_celula(ws, "Número", 2)) is None  # 1a linha de dado: sem zebra
    assert _fundo(_celula(ws, "Número", 3)) == COR_ZEBRA  # 2a linha: zebra
    assert _fundo(_celula(ws, "Emissor", 3)) == "FBECEB"  # mesma linha, mas confianca baixa VENCE a zebra
    assert _fundo(_celula(ws, "Número", 4)) is None  # 3a linha: sem zebra de novo
    assert _fundo(_celula(ws, "Número", 5)) == COR_ZEBRA  # 4a linha: zebra


def test_valor_total_do_resumo_em_destaque_negrito_maior_e_com_cor():
    wb, _ = _abrir(_doc(numero_documento="1", valor_total=229.0))
    celula = _celula(wb["Documentos"], "Valor total")
    assert celula.font.bold and celula.font.size == TAMANHO_FONTE_VALOR_TOTAL
    assert celula.font.color.rgb[-6:] == COR_DESTAQUE_VALOR_TOTAL
    assert not celula.font.italic


def test_valor_total_em_destaque_continua_com_fundo_e_italico_quando_corrigido():
    d = _doc(numero_documento="1", valor_total=229.0, corrigidos={"valor_total": 200.0})
    celula = _celula(_abrir(d)[0]["Documentos"], "Valor total")
    assert celula.font.bold and celula.font.italic  # destaque + marca de corrigido, os dois juntos
    assert _fundo(celula) == "E7F0EF"  # continua com o fundo de "corrigido"


def test_guia_de_todas_as_abas_colorida_igual_ao_cabecalho():
    wb, _ = _abrir(_doc(numero_documento="1"))
    for nome_aba in ["Relatório", *_ABAS_DE_DADOS]:
        assert wb[nome_aba].sheet_properties.tabColor.rgb[-6:] == COR_CABECALHO_FUNDO


def test_linha_de_total_soma_quantidade_e_valor_total_ignorando_texto():
    d = _doc(itens=[
        {"descricao": "a", "quantidade": 2.0, "valor_total": 10.0},
        {"descricao": "b", "quantidade": "1 un", "valor_total": 5.0},  # texto: SOMA ignora
    ])
    ws = _abrir(d)[0]["Itens"]
    linha_total = ws.max_row
    col_qtd, col_total = CABECALHOS_ITENS.index("Quantidade") + 1, CABECALHOS_ITENS.index("Valor total") + 1
    letra_qtd, letra_total = openpyxl.utils.get_column_letter(col_qtd), openpyxl.utils.get_column_letter(col_total)

    rotulo = ws.cell(linha_total, 1)
    assert rotulo.value == "Total" and rotulo.font.bold
    assert f"A{linha_total}:C{linha_total}" in [str(r) for r in ws.merged_cells.ranges]  # ID..Descricao mesclados
    assert ws.cell(linha_total, col_qtd).value == f"=SUM({letra_qtd}2:{letra_qtd}3)"
    total = ws.cell(linha_total, col_total)
    assert total.value == f"=SUM({letra_total}2:{letra_total}3)" and total.number_format == FORMATO_MOEDA
    # Valor unitario nao entra na soma (nao faz sentido somar preco unitario)
    col_unit = CABECALHOS_ITENS.index("Valor unitário") + 1
    assert ws.cell(linha_total, col_unit).value is None


def test_linha_de_total_ausente_quando_a_aba_itens_nao_tem_nenhum_item():
    ws = _abrir(_doc(numero_documento="1"))[0]["Itens"]
    assert ws.max_row == 1  # so cabecalho, sem linha de total


# ---------- total do lote (aba Documentos) ----------


def _lote_de_tres():
    return [
        _doc(arquivo="1.pdf", numero_documento="1", valor_total=100.0),
        _doc(arquivo="2.pdf", tipo_documento="boleto", numero_documento="2", valor_total=50.5),
        _doc(arquivo="3.pdf", numero_documento="3", valor_total="texto que nao virou numero"),
    ]


def test_total_do_lote_soma_o_valor_total_de_todos_os_documentos():
    ws = _abrir(*_lote_de_tres())[0]["Documentos"]
    col = CABECALHOS_RESUMO.index("Valor total") + 1
    letra = get_column_letter(col)
    linha_total = 5  # cabecalho + 3 documentos + esta
    assert ws.max_row == linha_total
    assert ws.cell(linha_total, col - 1).value == "Total do lote" and ws.cell(linha_total, col - 1).font.bold
    total = ws.cell(linha_total, col)
    assert total.value == f"=SUM({letra}2:{letra}4)"
    assert total.data_type == "f"  # formula de verdade (e a unica celula assim na aba)
    assert total.number_format == FORMATO_MOEDA and total.font.bold
    # o 3o documento tem valor_total em TEXTO: continua texto (nao vira 0), e a soma o ignora
    assert ws.cell(4, col).value == "texto que nao virou numero" and ws.cell(4, col).data_type == "s"


def test_total_do_lote_fica_fora_da_tabela_e_a_soma_cobre_exatamente_as_linhas_dela():
    ws = _abrir(*_lote_de_tres())[0]["Documentos"]
    ultima_coluna = get_column_letter(len(CABECALHOS_RESUMO))
    ref = ws.tables[NOMES_TABELA["Documentos"]].ref
    assert ref == f"A1:{ultima_coluna}4"  # a linha de total (5) nao entra na Tabela/filtro
    letra = get_column_letter(CABECALHOS_RESUMO.index("Valor total") + 1)
    ultima_linha_da_tabela = int(ref.rsplit(":", 1)[1].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    assert ws.cell(5, CABECALHOS_RESUMO.index("Valor total") + 1).value == f"=SUM({letra}2:{letra}{ultima_linha_da_tabela})"


def test_total_do_lote_ausente_com_um_unico_documento():
    ws = _abrir(_doc(numero_documento="1", valor_total=100.0))[0]["Documentos"]
    assert ws.max_row == 2  # cabecalho + o documento, sem linha de total
    assert all(c.data_type != "f" for linha in ws.iter_rows() for c in linha)
    assert all(c.value != "Total do lote" for linha in ws.iter_rows() for c in linha)


def test_total_do_lote_entra_na_area_de_impressao():
    ws = _abrir(*_lote_de_tres())[0]["Documentos"]
    assert ws.print_area == f"'Documentos'!$A$1:${get_column_letter(len(CABECALHOS_RESUMO))}$5"


def test_total_do_lote_nao_e_lido_como_documento_pela_tabela_nomeada():
    """Quem le pela Tabela (Power Query: Documentos[...]) so ve os documentos --
    a linha de total esta fora da ref. Quem le a aba crua ve uma linha a mais,
    com ID vazio (mesmo compromisso ja aceito na linha de total dos Itens)."""
    ws = _abrir(*_lote_de_tres())[0]["Documentos"]
    assert [r["ID"] for r in _linhas(ws)] == [1, 2, 3, None]
    assert [r["ID"] for r in _linhas_sem_total(ws)] == [1, 2, 3]


def test_nota_de_geracao_no_resumo_fora_da_tabela_e_do_filtro():
    buffer = gerar_excel([_doc(numero_documento="1")], data_geracao=date(2026, 9, 23))
    ws = openpyxl.load_workbook(buffer)["Documentos"]
    coluna_nota = len(CABECALHOS_RESUMO) + 2
    celula = ws.cell(1, coluna_nota)
    assert celula.value == "Gerado em 23/09/2026" and celula.font.italic
    # fora da Tabela e das colunas reais (ver docstring do modulo)
    ultima_coluna_tabela = openpyxl.utils.get_column_letter(len(CABECALHOS_RESUMO))
    assert ws.tables[NOMES_TABELA["Documentos"]].ref == f"A1:{ultima_coluna_tabela}2"


def test_nenhuma_aba_de_dados_tem_linha_de_titulo_mesclada_acima_do_cabecalho():
    """Decisao documentada no modulo: uma linha de titulo mesclada acima do
    cabecalho foi cogitada e DESCARTADA (testado com pandas.read_excel
    simulando um leitor automatico tipo Power Query -- ver o commit que
    trouxe a identidade visual) porque corrompe a leitura de quem assume
    "linha 1 = cabecalho". O cabecalho tem que continuar sendo a primeira
    linha de toda aba de DADOS, sem excecao. (A aba Relatorio fica de fora
    de proposito: nao e uma tabela, a linha 1 e o nome do aplicativo.)"""
    wb, _ = _abrir(_doc(numero_documento="1", itens=[{"descricao": "x"}],
                        campos_adicionais=[{"campo": "CFOP", "valor": "5102"}], avisos=["a"]))
    for nome_aba in _ABAS_DE_DADOS:
        ws = wb[nome_aba]
        cabecalhos_esperados = _CABECALHOS_POR_ABA[nome_aba]
        assert [ws.cell(1, c).value for c in range(1, len(cabecalhos_esperados) + 1)] == cabecalhos_esperados
        assert ws.merged_cells.ranges == [] or all(r.min_row > 1 for r in ws.merged_cells.ranges)


# ---------- acabamento senior (etapa 8) ----------


def test_metadados_do_arquivo_autor_e_o_aplicativo_nunca_uma_pessoa():
    buffer = gerar_excel([_doc(numero_documento="1")], data_geracao=date(2026, 9, 23))
    props = openpyxl.load_workbook(buffer).properties
    assert props.creator == NOME_APLICATIVO and props.lastModifiedBy == NOME_APLICATIVO
    assert NOME_APLICATIVO in props.title
    assert props.subject and "documentos" in props.subject.lower()
    assert props.created.date() == date(2026, 9, 23)


def test_impressao_area_cabecalho_repetido_paisagem_e_ajuste_de_escala():
    wb, _ = _abrir(_doc(numero_documento="1", itens=[{"descricao": "x", "valor_total": 1.0}]))
    for nome_aba in _ABAS_DE_DADOS:
        ws = wb[nome_aba]
        assert ws.print_title_rows == "$1:$1", f"{nome_aba}: cabecalho repetido"
        assert ws.page_setup.orientation == "landscape", f"{nome_aba}: paisagem"
        assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0, f"{nome_aba}: ajuste na largura"
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True, f"{nome_aba}: fitToPage precisa estar ligado"
    # area de impressao de Documentos NAO inclui a nota de geracao (2 colunas a mais)
    ultima_resumo = get_column_letter(len(CABECALHOS_RESUMO))
    assert wb["Documentos"].print_area == f"'Documentos'!$A$1:${ultima_resumo}$2"
    # a de Itens INCLUI a linha de total (e conteudo de leitura, so nao entra na Tabela/filtro)
    ws_itens = wb["Itens"]
    ultima_itens = get_column_letter(len(CABECALHOS_ITENS))
    assert ws_itens.print_area == f"'Itens'!$A$1:${ultima_itens}${ws_itens.max_row}"


def test_impressao_do_relatorio_e_retrato_sem_cabecalho_repetido():
    """Relatorio nao e tabela -- retrato (layout estreito), sem
    print_title_rows (nao ha cabecalho de coluna pra repetir), mas continua
    com ajuste de escala na largura, igual as outras abas."""
    ws = _abrir(_doc(numero_documento="1"))[0]["Relatório"]
    assert ws.page_setup.orientation == "portrait"
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert ws.print_title_rows is None


def test_documento_com_aviso_ganha_comentario_no_resumo_apontando_pra_aba_avisos():
    com_aviso = _doc(numero_documento="1", avisos=["a soma nao bate", "outro aviso"])
    sem_aviso = _doc(numero_documento="2")
    wb, _ = _abrir(com_aviso, sem_aviso)
    ws = wb["Documentos"]
    doc1 = _celula(ws, "Documento", 2)
    doc2 = _celula(ws, "Documento", 3)
    assert doc1.comment is not None and "2 avisos" in doc1.comment.text and "Avisos" in doc1.comment.text
    assert doc2.comment is None


def test_valor_total_dos_produtos_exposto_pela_danfe_vira_numero_no_excel():
    """Fecha o loop da coerencia Resumo x Itens: o valor que a DANFE ja
    comparava internamente com a soma dos itens agora tambem aparece, como
    NUMERO (nao so texto), na aba Campos adicionais -- comparavel de
    verdade, nao so por um aviso em texto."""
    d = _doc(campos_adicionais=[{"campo": "Valor Total dos Produtos", "valor": "215,03"}])
    celula = _celula(_abrir(d)[0]["Campos adicionais"], "Valor")
    assert celula.value == 215.03 and celula.number_format == FORMATO_MOEDA


def test_nenhum_dado_do_documento_vaza_pra_fora_das_abas():
    """A demo e publica e pode receber documento de qualquer pessoa -- o
    .xlsx nao pode guardar nada do documento fora do que aparece nas 4
    abas de dados. Confere no XML BRUTO (nao so pelo que o openpyxl expoe),
    porque metadado residual tipicamente NAO aparece navegando pela
    planilha normalmente.

    Achado ao investigar isto: o openpyxl carimba `modified` (docProps/
    core.xml) com o horario real do processo, sem forma de fixar isso antes
    do save() (ver comentario em gerar_excel) -- mas isso e o instante em
    que O SERVIDOR gerou o arquivo, nao um dado do documento ou de quem fez
    upload, entao nao entra nesta checagem."""
    sensivel = "nota_fiscal_do_FULANO_SECRETO_privado.pdf"
    d = _doc(
        arquivo=sensivel,
        numero_documento="000012345",
        emissor="FULANO SECRETO LTDA (CNPJ 11.222.333/0001-81)",
        destinatario="CICRANO PRIVADO (CPF 111.222.333-44)",
        valor_total=229.0,
        campos_adicionais=[{"campo": "Chave de Acesso", "valor": CHAVE_44}],
        confiancas={"emissor": "baixa"},
        corrigidos={"valor_total": 999.0},
    )
    buffer = gerar_excel([d])
    with zipfile.ZipFile(buffer) as zf:
        nomes = zf.namelist()
        assert "docProps/custom.xml" not in nomes  # nenhuma propriedade custom

        partes_com_dado = {n for n in nomes if n.startswith("xl/worksheets/sheet") or n == "xl/sharedStrings.xml"}
        partes_sem_dado = {n for n in nomes if n.endswith((".xml", ".rels")) and n not in partes_com_dado}
        assert partes_sem_dado  # a checagem abaixo nao pode ficar vazia por engano

        # a chave nunca deveria vazar nem no formato bruto nem no formatado (blocos de 4)
        termos_sensiveis = ["FULANO", "SECRETO", "CICRANO", sensivel, CHAVE_44, CHAVE_44_BLOCOS]
        for nome in partes_sem_dado:
            conteudo = zf.read(nome).decode("utf-8", errors="replace")
            for termo in termos_sensiveis:
                assert termo not in conteudo, f"{nome}: contem dado do documento ({termo!r}) fora das abas"

        # sanidade: os termos de fato aparecem em algum lugar (senao o teste seria vazio) --
        # a chave em si so aparece FORMATADA (Campos adicionais e Relatorio), nunca crua
        conteudo_dados = "".join(zf.read(n).decode("utf-8", errors="replace") for n in partes_com_dado)
        for termo in ["FULANO", "SECRETO", "CICRANO", sensivel, CHAVE_44_BLOCOS]:
            assert termo in conteudo_dados


def test_descricao_longa_nao_fixa_altura_de_linha_wrap_faz_o_resto():
    """O teto de largura (60) ja existia; o que garante que uma descricao de
    200 caracteres nao vira uma linha ilegivel e o Excel poder recalcular a
    altura sozinho ao abrir -- so acontece se a linha NAO tiver uma altura
    fixa (customHeight). So o cabecalho (linha 1) tem altura fixada."""
    d = _doc(itens=[{"descricao": "Produto " + "muito longo " * 20, "valor_total": 1.0}])
    ws = _abrir(d)[0]["Itens"]
    celula = _celula(ws, "Descrição")
    assert celula.alignment.wrap_text is True
    assert ws.column_dimensions["C"].width == LARGURA_MAXIMA
    linha_dado = ws.row_dimensions[2]
    assert linha_dado.height is None or linha_dado.customHeight in (None, False)


# ---------- Relatorio (etapa 9) ----------


def test_relatorio_e_a_aba_ativa_com_nome_do_sistema_e_data():
    buffer = gerar_excel([_doc(numero_documento="1")], data_geracao=date(2026, 9, 24))
    wb = openpyxl.load_workbook(buffer)
    assert wb.active.title == "Relatório"
    ws = wb["Relatório"]
    assert ws.cell(1, 2).value == NOME_APLICATIVO
    assert ws.cell(1, 2).font.bold and ws.cell(1, 2).font.size == TAMANHO_FONTE_APLICATIVO
    assert "24/09/2026" in ws.cell(2, 2).value


def test_relatorio_titulo_do_documento_com_arquivo_modo_e_origem():
    d = _doc(arquivo="nota.pdf", numero_documento="000012345")
    ws = _abrir(d)[0]["Relatório"]
    linha = _linha_com_texto(ws, "Nota fiscal 000012345")
    assert linha is not None
    subtitulo = ws.cell(linha + 1, 2).value
    assert "nota.pdf" in subtitulo and "Modo básico" in subtitulo and "Texto digital" in subtitulo


def test_relatorio_campo_obrigatorio_vazio_mostra_nao_encontrado():
    d = _doc(tipo_documento="boleto", numero_documento=None, data_vencimento=None)  # os dois obrigatorios p/ boleto
    ws = _abrir(d)[0]["Relatório"]
    for rotulo in ("Número", "Data de vencimento"):
        celula = _valor_por_rotulo(ws, rotulo)
        assert celula.value == "não encontrado"
        assert celula.font.italic and celula.font.color.rgb[-6:] == COR_AUSENTE


def test_relatorio_campo_opcional_vazio_some_a_linha_inteira():
    d = _doc(tipo_documento="relatorio", numero_documento=None)  # tipo "relatorio": nada e obrigatorio
    ws = _abrir(d)[0]["Relatório"]
    assert _linha_com_texto(ws, "Número") is None


def test_relatorio_destaque_de_confianca_com_fundo_e_comentario_no_valor():
    d = _doc(numero_documento="1", emissor="E (CNPJ 11.222.333/0001-81)", confiancas={"emissor": "baixa"})
    ws = _abrir(d)[0]["Relatório"]
    linha_secao = _linha_com_texto(ws, "EMISSOR")
    celula = _valor_por_rotulo(ws, "Nome", a_partir_de=linha_secao)
    assert _fundo(celula) == "FBECEB" and celula.comment is not None and "baixa" in celula.comment.text


def test_relatorio_valor_total_em_fonte_grande_na_cor_de_destaque():
    ws = _abrir(_doc(numero_documento="1", valor_total=229.0))[0]["Relatório"]
    celula = _valor_por_rotulo(ws, "Valor total")
    assert celula.value == 229.0 and celula.number_format == FORMATO_MOEDA
    assert celula.font.bold and celula.font.size == TAMANHO_FONTE_VALOR_PRINCIPAL
    assert celula.font.color.rgb[-6:] == COR_DESTAQUE_VALOR_TOTAL


def test_relatorio_chave_de_acesso_em_blocos_de_4_digitos():
    d = _doc(campos_adicionais=[{"campo": "Chave de Acesso", "valor": CHAVE_44}])
    ws = _abrir(d)[0]["Relatório"]
    celula = _valor_por_rotulo(ws, "Chave de Acesso")
    assert celula.value == CHAVE_44_BLOCOS


def test_relatorio_valor_total_tem_linha_alta_o_bastante_pra_fonte_grande():
    """Bug real: o Excel NAO recalcula a altura da linha sozinho quando a
    celula com a fonte maior esta mesclada (limitacao do proprio Excel) --
    sem altura explicita, "R$ 229,00" em fonte 20 aparecia cortado pela
    metade."""
    ws = _abrir(_doc(numero_documento="1", valor_total=229.0))[0]["Relatório"]
    linha = _linha_com_texto(ws, "Valor total")
    altura = ws.row_dimensions[linha].height
    assert altura is not None and altura >= TAMANHO_FONTE_VALOR_PRINCIPAL * 1.2


def test_relatorio_rodape_da_legenda_tem_espaco_antes_do_ultimo_bloco():
    """A propria linha imediatamente acima do rodape tem que ser um espacador
    DEDICADO da legenda (altura fixa, sem conteudo) -- nao basta o espaco que
    ja sobra depois do ultimo bloco do documento (que existe independente
    disso e nao e garantido em todo layout, ex: sem "Dados adicionais")."""
    d = _doc(numero_documento="1", emissor="E (CNPJ 11.222.333/0001-81)", confiancas={"emissor": "baixa"})
    ws = _abrir(d)[0]["Relatório"]
    texto_baixa = next(t for e, t in LEGENDA_LINHAS if e == "baixa")
    linha_legenda = _linha_com_texto(ws, texto_baixa, coluna=3)
    linha_espacador = linha_legenda - 1
    assert all(ws.cell(linha_espacador, c).value in (None, "") for c in range(1, 7))
    assert ws.row_dimensions[linha_espacador].height == 12


def test_relatorio_campo_monetario_adicional_fica_em_valores_o_resto_em_dados_adicionais():
    d = _doc(campos_adicionais=[
        {"campo": "Valor do Documento", "valor": "1.000,00"},
        {"campo": "CFOP", "valor": "5102"},
    ])
    ws = _abrir(d)[0]["Relatório"]
    linha_valores = _linha_com_texto(ws, "VALORES")
    linha_dados = _linha_com_texto(ws, "DADOS ADICIONAIS")
    assert linha_valores < _linha_com_texto(ws, "Valor do Documento") < linha_dados
    assert _linha_com_texto(ws, "CFOP") > linha_dados
    assert _valor_por_rotulo(ws, "Valor do Documento").value == 1000.0


def test_relatorio_tabela_de_itens_compacta_sem_ser_tabela_do_excel():
    d = _doc(itens=[{"descricao": "Mochila", "quantidade": 1.0, "valor_unitario": 215.03, "valor_total": 215.03}])
    ws = _abrir(d)[0]["Relatório"]
    linha = _linha_com_texto(ws, "Mochila")
    assert linha is not None
    assert ws.cell(linha, 6).value == 215.03 and ws.cell(linha, 6).number_format == FORMATO_MOEDA
    assert ws.tables == {}  # formatada em celulas, NAO uma Tabela nomeada (plano aprovado)


def test_relatorio_sem_itens_nao_mostra_a_secao():
    ws = _abrir(_doc(numero_documento="1"))[0]["Relatório"]
    assert _linha_com_texto(ws, "ITENS") is None


def test_relatorio_pontos_de_atencao_so_aparece_quando_ha_aviso():
    com_aviso = _abrir(_doc(numero_documento="1", avisos=["a soma não bate"]))[0]["Relatório"]
    linha = _linha_com_texto(com_aviso, "PONTOS DE ATENÇÃO")
    assert linha is not None and "a soma não bate" in com_aviso.cell(linha + 1, 2).value

    sem_aviso = _abrir(_doc(numero_documento="1"))[0]["Relatório"]
    assert _linha_com_texto(sem_aviso, "PONTOS DE ATENÇÃO") is None


def test_relatorio_quebra_de_pagina_entre_documentos_do_lote_mas_nao_com_1_so():
    ws_lote = _abrir(_doc(arquivo="1.pdf", numero_documento="1"), _doc(arquivo="2.pdf", numero_documento="2"))[0]["Relatório"]
    assert len(ws_lote.row_breaks.brk) == 1  # 2 documentos -> 1 quebra, entre os dois

    ws_um = _abrir(_doc(numero_documento="1"))[0]["Relatório"]
    assert len(ws_um.row_breaks.brk) == 0


def test_relatorio_rodape_so_mostra_as_cores_que_aparecem_de_verdade():
    texto_baixa = next(t for e, t in LEGENDA_LINHAS if e == "baixa")
    com_baixa = _abrir(_doc(numero_documento="1", emissor="E (CNPJ 11.222.333/0001-81)", confiancas={"emissor": "baixa"}))[0]["Relatório"]
    linha = _linha_com_texto(com_baixa, texto_baixa, coluna=3)
    assert linha is not None and _fundo(com_baixa.cell(linha, 2)) == FUNDOS["baixa"]
    # as OUTRAS legendas (media/corrigido) nao aparecem -- so a que foi usada
    for estado, texto in LEGENDA_LINHAS:
        if estado != "baixa":
            assert _linha_com_texto(com_baixa, texto, coluna=3) is None

    sem_nenhuma = _abrir(_doc(numero_documento="1"))[0]["Relatório"]
    for _, texto in LEGENDA_LINHAS:
        assert _linha_com_texto(sem_nenhuma, texto, coluna=3) is None


def test_relatorio_sem_grade_e_com_margem_estreita():
    ws = _abrir(_doc(numero_documento="1"))[0]["Relatório"]
    assert ws.sheet_view.showGridLines is False
    assert ws.column_dimensions["A"].width < ws.column_dimensions["B"].width


# ---------- seguranca e robustez ----------


def test_texto_que_parece_formula_fica_texto_e_nao_vira_formula():
    """O openpyxl trata string com "=" no inicio como formula; o texto vem de
    PDF (nao confiavel), entao um emissor "=HYPERLINK(...)" nao pode executar."""
    perigoso = '=HYPERLINK("http://x","clique")'
    d = _doc(emissor=perigoso, destinatario="=1+1", numero_documento="=2+2",
             itens=[{"descricao": "=SUM(A1:A2)"}], campos_adicionais=[{"campo": "Nota", "valor": "=A1"}],
             avisos=["=cmd|' /C calc'!A0"])
    wb, buffer = _abrir(d)

    assert _celula(wb["Documentos"], "Emissor").value == perigoso
    assert _celula(wb["Documentos"], "Destinatário").value == "=1+1"
    assert _celula(wb["Itens"], "Descrição").value == "=SUM(A1:A2)"

    # a UNICA formula do arquivo e a soma da linha de total dos Itens -- gerada
    # pelo nosso proprio codigo, com referencia de celula fixa, nunca a partir
    # de texto do PDF (ver test_linha_de_total_soma_quantidade_e_valor_total)
    ws_itens = wb["Itens"]
    linha_total = ws_itens.max_row
    for ws in wb:
        ultima_linha_dado = ws.max_row - 1 if ws is ws_itens else ws.max_row
        for linha in ws.iter_rows(max_row=ultima_linha_dado):
            assert all(c.data_type != "f" for c in linha), f"{ws.title}: celula virou formula (fora da linha de total)"
    col_total = CABECALHOS_ITENS.index("Valor total") + 1
    assert ws_itens.cell(linha_total, col_total).data_type == "f"
    assert ws_itens.cell(linha_total, col_total).value == "=SUM(F2:F2)"

    buffer.seek(0)
    # no XML inteiro, as UNICAS formulas sao as 2 da linha de total dos Itens
    # (Quantidade + Valor total) -- nenhum texto perigoso do PDF escapou como <f>
    n_formulas = 0
    with zipfile.ZipFile(buffer) as zf:
        for nome in zf.namelist():
            if nome.startswith("xl/worksheets/sheet"):
                xml = zf.read(nome).decode("utf-8")
                n_formulas += xml.count("<f>") + xml.count("<f ")
    assert n_formulas == 2


def test_caracteres_de_controle_sao_removidos_em_vez_de_quebrar_o_arquivo():
    d = _doc(emissor="EMPRESA\x00\x01\x0b LTDA", numero_documento="12\x1f3")
    ws = _abrir(d)[0]["Documentos"]
    assert _celula(ws, "Emissor").value == "EMPRESA LTDA"
    assert _celula(ws, "Número").value == "123"


def test_texto_gigante_e_truncado_no_limite_do_excel():
    ws = _abrir(_doc(emissor="X" * 40000))[0]["Documentos"]
    assert len(_celula(ws, "Emissor").value) == 32767


def test_larguras_ajustadas_ao_conteudo_com_teto():
    d = _doc(numero_documento="1", emissor="E" * 200, destinatario="Curto",
             campos_adicionais=[{"campo": "Chave de Acesso", "valor": CHAVE_44}])
    wb, _ = _abrir(d)
    resumo = wb["Documentos"]
    largura = lambda ws, letra: ws.column_dimensions[letra].width
    assert largura(resumo, "H") == 60  # Emissor: 200 chars -> teto
    assert largura(resumo, "C") >= len("Nota fiscal 1")  # Documento cabe
    assert largura(resumo, "A") <= 10  # ID: estreita
    assert largura(wb["Campos adicionais"], "D") >= len(CHAVE_44)  # a chave inteira cabe (44 < 60)


# ---------- endpoint ----------


def _corpo(resposta):
    async def ler():
        return b"".join([pedaco async for pedaco in resposta.body_iterator])

    return asyncio.run(ler())


def test_endpoint_devolve_xlsx_com_nome_de_um_documento_ou_de_lote():
    um = asyncio.run(main.export_excel(ExportarExcelRequest(documentos=[_doc(numero_documento="1")])))
    assert um.headers["content-disposition"] == 'attachment; filename="documento_extraido.xlsx"'
    assert openpyxl.load_workbook(BytesIO(_corpo(um))).sheetnames == _SHEETNAMES_ESPERADOS

    varios = asyncio.run(main.export_excel(ExportarExcelRequest(documentos=[_doc(), _doc()])))
    assert varios.headers["content-disposition"] == 'attachment; filename="documentos_extraidos.xlsx"'


def test_requisicao_sem_documentos_e_rejeitada():
    with pytest.raises(ValidationError):
        ExportarExcelRequest(documentos=[])
